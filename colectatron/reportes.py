"""Tablas de resultados y el planner en Excel."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ConfigColecta
from .formulario import Voluntario


def rango_horario(bloques: list[str]) -> str:
    """Une bloques seguidos: ["10:00 - 11:30", "11:30 - 13:00"] → "10:00 - 13:00"."""
    if not bloques:
        return ""
    inicio = bloques[0].split("-")[0].strip()
    fin = bloques[-1].split("-")[-1].strip()
    return f"{inicio} - {fin}"


def por_persona(config: ConfigColecta, voluntarios: list[Voluntario], asignaciones: pd.DataFrame) -> pd.DataFrame:
    """Una fila por persona y día: dónde y a qué hora le toca. Útil para avisar por WhatsApp."""
    columnas = ["Nombre", "Día", "Lugar", "Horario", "Bloques", "Auto", "Teléfono", "Correo"]
    if asignaciones.empty:
        return pd.DataFrame(columns=columnas)
    datos = {v.id: v for v in voluntarios}
    orden = {s: i for i, s in enumerate(config.bloques)}
    filas = []
    for (vid, d, lugar), grupo in asignaciones.groupby(["id", "dia", "lugar"], sort=False):
        v = datos[vid]
        bloques = sorted(grupo["bloque"], key=orden.get)
        filas.append(
            {
                "Nombre": v.nombre,
                "Día": d,
                "Lugar": lugar,
                "Horario": rango_horario(bloques),
                "Bloques": len(bloques),
                "Auto": "Sí" if d in v.auto else "",
                "Teléfono": v.telefono,
                "Correo": v.correo,
            }
        )
    df = pd.DataFrame(filas, columns=columnas)
    orden_dia = {d: i for i, d in enumerate(config.dias)}
    return df.sort_values(["Nombre", "Día"], key=lambda s: s.map(orden_dia) if s.name == "Día" else s).reset_index(
        drop=True
    )


def sin_turno(voluntarios: list[Voluntario], asignaciones: pd.DataFrame) -> pd.DataFrame:
    con_turno = set(asignaciones["id"]) if not asignaciones.empty else set()
    filas = [
        {
            "Nombre": v.nombre,
            "Rol": v.rol,
            "Motivo": "Disponible, pero no cupo" if v.tiene_disponibilidad else "No marcó horarios",
            "Teléfono": v.telefono,
            "Correo": v.correo,
        }
        for v in voluntarios
        if v.id not in con_turno
    ]
    return pd.DataFrame(filas, columns=["Nombre", "Rol", "Motivo", "Teléfono", "Correo"]).sort_values("Nombre")


def cobertura(config: ConfigColecta, asignaciones: pd.DataFrame) -> pd.DataFrame:
    """Personas por (día, lugar, bloque), incluyendo los bloques vacíos."""
    indice = pd.MultiIndex.from_product(
        [config.dias, config.nombres_lugares, config.bloques], names=["dia", "lugar", "bloque"]
    )
    if asignaciones.empty:
        conteo = pd.Series(0, index=indice)
    else:
        conteo = asignaciones.groupby(["dia", "lugar", "bloque"]).size().reindex(indice, fill_value=0)
    return conteo.rename("personas").reset_index()


def resumen(config: ConfigColecta, voluntarios: list[Voluntario], asignaciones: pd.DataFrame) -> dict[str, float]:
    cob = cobertura(config, asignaciones)
    con_turno = asignaciones["id"].nunique() if not asignaciones.empty else 0
    return {
        "personas": len(voluntarios),
        "con_turno": con_turno,
        "porcentaje_con_turno": 100 * con_turno / len(voluntarios) if voluntarios else 0,
        "turnos": len(asignaciones),
        "bloques_cubiertos": int((cob["personas"] > 0).sum()),
        "bloques_totales": len(cob),
    }


def _celda_persona(v: Voluntario, dia: str) -> str:
    marcas = ("⭐ " if v.es_jefe else "") + ("🚗 " if dia in v.auto else "")
    lineas = [f"{marcas}{v.nombre}"]
    if v.telefono:
        lineas.append(f"📞 {v.telefono}")
    return "\n".join(lineas)


def _ajustar_hoja(ws, ancho: int = 32, envolver: bool = True) -> None:
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho if i > 1 else 16
    for fila in ws.iter_rows():
        for celda in fila:
            celda.alignment = Alignment(wrap_text=envolver, vertical="top")
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "B2"


def _nombre_hoja(texto: str, usadas: set[str]) -> str:
    base = "".join(c for c in texto if c not in "[]:*?/\\")[:31] or "Hoja"
    nombre, i = base, 2
    while nombre in usadas:
        nombre = f"{base[:28]} {i}"
        i += 1
    usadas.add(nombre)
    return nombre


def planner_excel(
    config: ConfigColecta,
    voluntarios: list[Voluntario],
    asignaciones: pd.DataFrame,
    bodegas: pd.DataFrame | None = None,
    avisos: list[str] | None = None,
) -> bytes:
    """Excel con una hoja por día (bloques × lugares) y hojas de apoyo."""
    datos = {v.id: v for v in voluntarios}
    salida = io.BytesIO()
    usadas: set[str] = set()
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        for d in config.dias:
            tabla = pd.DataFrame("", index=config.bloques, columns=config.nombres_lugares)
            if not asignaciones.empty:
                del_dia = asignaciones[asignaciones["dia"] == d]
                for (s, lugar), grupo in del_dia.groupby(["bloque", "lugar"]):
                    tabla.loc[s, lugar] = "\n\n".join(_celda_persona(datos[i], d) for i in grupo["id"])
            tabla.index.name = "Bloque"
            hoja = _nombre_hoja(d, usadas)
            tabla.to_excel(writer, sheet_name=hoja)
            _ajustar_hoja(writer.sheets[hoja])

        hojas = [
            ("Por persona", por_persona(config, voluntarios, asignaciones)),
            ("Sin turno", sin_turno(voluntarios, asignaciones)),
        ]
        if bodegas is not None and not bodegas.empty:
            hojas.append(("Casas-bodega", bodegas))
        if avisos:
            hojas.append(("Avisos", pd.DataFrame({"Aviso": avisos})))
        for nombre, df in hojas:
            hoja = _nombre_hoja(nombre, usadas)
            df.to_excel(writer, sheet_name=hoja, index=False)
            _ajustar_hoja(writer.sheets[hoja], ancho=24, envolver=False)
            writer.sheets[hoja].freeze_panes = "A2"

        leyenda = pd.DataFrame({"Símbolo": ["⭐", "🚗"], "Significado": ["Jefe/a", "Tiene auto ese día"]})
        hoja = _nombre_hoja("Leyenda", usadas)
        leyenda.to_excel(writer, sheet_name=hoja, index=False)
    return salida.getvalue()
