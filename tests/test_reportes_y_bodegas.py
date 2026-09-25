from __future__ import annotations

import io
from dataclasses import replace

import openpyxl
from conftest import persona

from colectatron.bodegas import asignar_bodegas
from colectatron.config import ConfigColecta, Reglas, config_insumos
from colectatron.ejemplo import respuestas_ejemplo
from colectatron.formulario import leer_respuestas
from colectatron.plantilla_form import script_apps
from colectatron.reportes import planner_excel, por_persona, rango_horario, resumen, sin_turno
from colectatron.solver import resolver


def test_rango_horario():
    assert rango_horario(["10:00 - 11:30", "11:30 - 13:00"]) == "10:00 - 13:00"
    assert rango_horario([]) == ""


def test_bodegas_prefiere_misma_comuna(config_chica):
    config = replace(config_chica, reglas=Reglas(nadie_solo=False))
    vols = [
        persona("trabaja", lugares=["A"]),
        persona("casa_lejos", bloques=[], bodega=True, comuna="Maipú"),
        persona("casa_cerca", bloques=[], bodega=True, comuna="Ñuñoa"),
    ]
    r = resolver(config, vols)
    bodegas, avisos = asignar_bodegas(config, vols, r.asignaciones)
    assert bodegas.iloc[0]["Casa-bodega"] == "casa_cerca"
    assert not avisos

    # Con tiempos de traslado manda la cercanía real.
    minutos = {("A", "casa_lejos"): 5, ("A", "casa_cerca"): 30}
    bodegas, _ = asignar_bodegas(config, vols, r.asignaciones, minutos)
    assert bodegas.iloc[0]["Casa-bodega"] == "casa_lejos"


def test_bodegas_avisa_si_no_hay(config_chica):
    vols = [persona("trabaja", lugares=["A"])]
    r = resolver(config_chica, vols)
    bodegas, avisos = asignar_bodegas(config_chica, vols, r.asignaciones)
    assert avisos and "Sin casa-bodega" in bodegas.iloc[0]["Casa-bodega"]


def test_flujo_completo_y_planner():
    config = config_insumos()
    lectura = leer_respuestas(respuestas_ejemplo(config, 40), config)
    r = resolver(config, lectura.voluntarios)
    assert r.optimo and not r.vacio

    res = resumen(config, lectura.voluntarios, r.asignaciones)
    assert res["con_turno"] + len(sin_turno(lectura.voluntarios, r.asignaciones)) == res["personas"]

    tabla = por_persona(config, lectura.voluntarios, r.asignaciones)
    assert tabla["Bloques"].sum() == len(r.asignaciones)

    bodegas, _ = asignar_bodegas(config, lectura.voluntarios, r.asignaciones)
    datos = planner_excel(config, lectura.voluntarios, r.asignaciones, bodegas, r.avisos)
    libro = openpyxl.load_workbook(io.BytesIO(datos))
    assert libro.sheetnames[:2] == config.dias
    assert {"Por persona", "Sin turno", "Casas-bodega", "Leyenda"} <= set(libro.sheetnames)
    hoja = libro[config.dias[0]]
    assert [c.value for c in hoja[1]][1:] == config.nombres_lugares


def test_config_json_ida_y_vuelta():
    config = config_insumos()
    assert ConfigColecta.from_json(config.to_json()) == config
    assert config.validar() == []
    assert replace(config, dias=[]).validar()


def test_script_del_form_incluye_configuracion():
    config = config_insumos()
    codigo = script_apps(config)
    assert "function crearFormulario" in codigo
    for texto in config.dias + config.nombres_lugares + config.bloques:
        assert texto in codigo
